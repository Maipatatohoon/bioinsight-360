from pptx import Presentation

prs = Presentation('/mnt/hdd/Downloads/BioInsight360_Frontend_Progress_Report.pptx')

for i, slide in enumerate(prs.slides):
    print(f"--- Slide {i+1} ---")
    for shape in slide.shapes:
        if hasattr(shape, "text"):
            print(shape.text)
